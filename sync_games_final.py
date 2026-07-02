#!/usr/bin/env python
"""Final sync: add missing game 1416 to JSON, remove None games."""
import json
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

print("=" * 70)
print("FINAL SYNC: ADD MISSING GAME 1416")
print("=" * 70)

# Load Excel
excel_path = Path("games.xlsx")
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# Load JSON
json_path = Path("data/all_games.json")
with open(json_path, 'r', encoding='utf-8') as f:
    json_games = json.load(f)

print(f"\n1. Current state:")
print(f"   JSON games: {len(json_games)}")

# Remove None/invalid games
json_games_clean = [g for g in json_games if g.get('gameNumber')]
print(f"   After removing invalid: {len(json_games_clean)}")

# Get game 1416 from Excel
game_1416 = None
for row_idx, row in enumerate(ws.iter_rows(min_row=4, max_row=100), start=4):
    game_num = row[1].value  # Column B = gameNumber
    if str(game_num) == '1416':
        # Extract all fields
        date_val = row[0].value
        if isinstance(date_val, datetime):
            date_val = date_val.strftime('%Y-%m-%d')

        game_1416 = {
            'season': '2025/2026',
            'date': str(date_val),
            'gameNumber': str(game_num),
            'location': str(row[2].value) if row[2].value else '',
            'transportation': float(row[3].value) if row[3].value else 0,
            'food': float(row[4].value) if row[4].value else 0,
            'gamePayment': float(row[5].value) if row[5].value else 0,
            'paidStatus': str(row[6].value) if row[6].value else '',
            'paymentDate': None,
            'observations': None,
        }
        if row[7].value:
            payment_date = row[7].value
            if isinstance(payment_date, datetime):
                payment_date = payment_date.strftime('%Y-%m-%d')
            game_1416['paymentDate'] = payment_date

        print(f"\n2. Found game 1416 in Excel:")
        print(f"   {game_1416}")
        break

if game_1416:
    # Check if 1416 already exists in JSON
    exists = any(str(g.get('gameNumber')) == '1416' for g in json_games_clean)

    if not exists:
        print(f"\n3. Adding game 1416 to JSON...")
        json_games_clean.append(game_1416)
        print(f"   [OK] Added game 1416")
    else:
        print(f"\n3. Game 1416 already exists in JSON")

    # Save JSON
    print(f"\n4. Saving JSON file...")
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(json_games_clean, f, indent=2, ensure_ascii=False)
    print(f"   [OK] Saved {len(json_games_clean)} games")

    # Final verification
    print(f"\n5. Final verification...")

    # Get Excel game numbers
    excel_games = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=4, max_row=100), start=4):
        game_num = row[1].value
        if game_num:
            excel_games.append(str(game_num))

    excel_set = set(excel_games)
    json_set = set(str(g.get('gameNumber')) for g in json_games_clean)

    missing = excel_set - json_set
    extra = json_set - excel_set

    print(f"   Excel games: {len(excel_games)}")
    print(f"   JSON games: {len(json_games_clean)}")
    print(f"   Missing in JSON: {missing}")
    print(f"   Extra in JSON: {extra}")

    if not missing and not extra:
        print(f"\n[SUCCESS] Excel and JSON are now in sync!")
        print(f"All {len(json_games_clean)} games are properly stored.")
    else:
        print(f"\n[WARNING] Still some mismatches")
else:
    print("[ERROR] Could not find game 1416 in Excel")

print("\n" + "=" * 70)
