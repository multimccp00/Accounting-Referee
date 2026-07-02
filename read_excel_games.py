#!/usr/bin/env python
"""Read games from Excel file and compare with database/JSON."""
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

# Load Excel
excel_path = Path("games.xlsx")
if not excel_path.exists():
    print("Excel file not found")
    sys.exit(1)

print("Reading Excel file...")
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

print(f"Sheet name: {ws.title}")
print(f"Dimensions: {ws.dimensions}\n")

# Read headers
headers = []
for cell in ws[1]:
    if cell.value:
        headers.append(cell.value)

print(f"Headers: {headers}\n")

# Read data
games = []
for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if not row[0]:  # Stop at empty rows
        break

    # Map to game structure
    game = {}
    for col_idx, header in enumerate(headers):
        if col_idx < len(row):
            game[header] = row[col_idx]

    games.append(game)

    # Show first few games
    if row_idx <= 5:
        print(f"Row {row_idx}: {game}")

print(f"\nTotal games in Excel: {len(games)}")

# Check current database
print("\n" + "=" * 70)
print("Checking current database...")

json_path = Path("data/all_games.json")
if json_path.exists():
    with open(json_path, 'r', encoding='utf-8') as f:
        current_data = json.load(f)

    current_games = current_data.get('games', [])
    print(f"Games in JSON: {len(current_games)}")

    if current_games:
        print(f"First game: {current_games[0]}")
else:
    print("JSON file not found")
    current_games = []

# Identify missing games
print("\n" + "=" * 70)
print("Comparing Excel vs JSON...")

# Create set of existing game identifiers
existing_ids = set()
for game in current_games:
    # Try to create a unique identifier
    key = (game.get('date'), game.get('team'), game.get('opponent'))
    existing_ids.add(key)

print(f"Existing unique game identifiers: {len(existing_ids)}")

# Find missing games
missing_games = []
for game in games:
    # Find equivalent in Excel
    excel_key = (game.get('Date') or game.get('date'),
                 game.get('Team') or game.get('team'),
                 game.get('Opponent') or game.get('opponent'))

    if excel_key not in existing_ids:
        missing_games.append(game)

print(f"Missing games found: {len(missing_games)}")

if missing_games:
    print("\nMissing games:")
    for i, game in enumerate(missing_games[:10], 1):
        print(f"  {i}. {game}")
    if len(missing_games) > 10:
        print(f"  ... and {len(missing_games) - 10} more")

print("\n" + "=" * 70)
