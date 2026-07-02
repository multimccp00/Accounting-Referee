#!/usr/bin/env python
"""Find the correct game numbers - row 31 is 1417, not 1416."""
import sys
from pathlib import Path

try:
    import openpyxl
except:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

excel_path = Path("games.xlsx")
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

print("FINDING DUPLICATE 1417")
print("=" * 100)

# Find all rows with 1417
rows_with_1417 = []
for row_num in range(4, 61):
    cell_value = ws[row_num][1].value  # Column B
    if cell_value and str(cell_value).strip() == '1417':
        # Get all data for this row
        row = ws[row_num]
        date_val = row[0].value
        location_val = row[2].value
        rows_with_1417.append({
            'row': row_num,
            'date': date_val,
            'location': location_val,
            'gameNum': '1417'
        })

print(f"Found {len(rows_with_1417)} rows with game 1417:\n")
for r in rows_with_1417:
    print(f"  Row {r['row']}: {r['date']} - {r['location']}")

if len(rows_with_1417) == 2:
    print(f"\nOne of these is a duplicate. Looking for the missing game number...")

    # Load all game numbers from Excel
    all_nums = []
    for row_num in range(4, 61):
        cell_val = ws[row_num][1].value
        if cell_val and str(cell_val).strip() not in ['Totals', 'gameNumber']:
            all_nums.append(str(cell_val).strip())

    print(f"All game numbers in Excel: {sorted(set(all_nums))}")

    # Find which numbers are missing in sequence
    print("\nLooking for missing game number near 1417...")
    for test_num in [1416, 1418, 1419, 1421, 1422, 1423]:
        if str(test_num) not in all_nums:
            print(f"  {test_num} - MISSING (could be what 1417 should be)")
